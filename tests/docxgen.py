"""Word documents for the tests, written by hand (every value invented): a name cut between
two runs of different formatting, a title, a list, a table typed by its headers, a header
and a footer, a footnote, a comment, a text box, a tracked deletion, an email link, the
document properties, a thumbnail and an image. Usage: from docxgen import contract"""
import io, zipfile

NS = ('xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
      'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
      'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
      'xmlns:w14="http://schemas.microsoft.com/office/word/2010/wordml" '
      'xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
      'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
      'xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape" '
      'mc:Ignorable="w14"')
HEAD = '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
B, I = "<w:rPr><w:b/></w:rPr>", "<w:rPr><w:i/></w:rPr>"


def run(text, props=""):
    space = ' xml:space="preserve"' if text != text.strip() else ""
    return f"<w:r>{props}<w:t{space}>{text}</w:t></w:r>"


def para(*runs, style=None, num=False):
    ppr = ""
    if style or num:
        ppr = "<w:pPr>" + (f'<w:pStyle w:val="{style}"/>' if style else "") + \
              ('<w:numPr><w:ilvl w:val="0"/><w:numId w:val="1"/></w:numPr>' if num else "") + \
              '<w:tabs><w:tab w:val="left" w:pos="720"/></w:tabs></w:pPr>'
    return f"<w:p>{ppr}{''.join(runs)}</w:p>"


def cell(text, props=""):
    return f'<w:tc><w:tcPr><w:tcW w:w="3000" w:type="dxa"/></w:tcPr>{para(run(text, props))}</w:tc>'


def table(rows):
    return ('<w:tbl><w:tblPr><w:tblStyle w:val="TableGrid"/></w:tblPr><w:tblGrid><w:gridCol w:w="3000"/>'
            '<w:gridCol w:w="3000"/><w:gridCol w:w="3000"/></w:tblGrid>'
            + "".join("<w:tr>" + "".join(cell(c, B if i == 0 else "") for c in r) + "</w:tr>"
                      for i, r in enumerate(rows)) + "</w:tbl>")


TEXTBOX = ('<w:r><w:drawing><wp:anchor distT="0" distB="0" distL="0" distR="0" simplePos="0" relativeHeight="1" '
           'behindDoc="0" locked="0" layoutInCell="1" allowOverlap="1"><wp:simplePos x="0" y="0"/>'
           '<wp:positionH relativeFrom="column"><wp:posOffset>0</wp:posOffset></wp:positionH>'
           '<wp:positionV relativeFrom="paragraph"><wp:posOffset>0</wp:posOffset></wp:positionV>'
           '<wp:extent cx="2000000" cy="500000"/><wp:wrapNone/><wp:docPr id="1" name="Zone de texte 1"/>'
           '<a:graphic><a:graphicData uri="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">'
           '<wps:wsp><wps:spPr/><wps:txbx><w:txbxContent>'
           + para(run("Témoin : "), run("Sophie Lambert", B)) +
           '</w:txbxContent></wps:txbx><wps:bodyPr/></wps:wsp></a:graphicData></a:graphic></wp:anchor></w:drawing></w:r>')


def contract():
    body = "".join([
        para(run("Affaire Julien Moreau contre Karim Hamidi"), style="Heading1"),
        # the name cut in the middle, second half in italics, first in bold
        para(run("Maître Claire Dubreuil représente "), run("M. Julien Mo", B), run("reau", I),
             run(", né le 14 août 1979 à Lyon.")),
        para(run("Objet :"), "<w:r><w:tab/></w:r>", run("litige avec Karim Hamidi")),
        para(run("Adresse : 14 rue des Lilas"), "<w:r><w:br/></w:r>", run("69003 Lyon")),
        para(run("Premier point : le contrat signé par Julien Moreau."), num=True),
        para(run("Second point : la lettre de Karim "), run("Hamidi", B), run(" du 3 mars."), num=True),
        table([["Nom", "Prénom", "Date de naissance"],
               ["MOREAU", "Julien", "14/08/1979"],
               ["HAMIDI", "Karim", "02/11/1985"]]),
        para(run("Témoin cité en annexe."), TEXTBOX),
        para(run("Le client "), '<w:del w:id="1" w:author="Claire Dubreuil" w:date="2026-09-01T10:00:00Z">'
             '<w:r><w:delText>Julien Moreau</w:delText></w:r></w:del>',
             '<w:ins w:id="2" w:author="Claire Dubreuil" w:date="2026-09-01T10:00:00Z">'
             + run("M. Moreau") + '</w:ins>', run(" confirme.")),
        para(run("Écrire à "), '<w:hyperlink r:id="rIdMail">' + run("julien.moreau@exemple.fr",
             '<w:rPr><w:rStyle w:val="Hyperlink"/></w:rPr>') + "</w:hyperlink>",
             run(" avant le 3 octobre."),
             '<w:r><w:footnoteReference w:id="1"/></w:r>',
             '<w:commentRangeStart w:id="0"/>', run(" Fin."), '<w:commentRangeEnd w:id="0"/>',
             '<w:r><w:commentReference w:id="0"/></w:r>'),
        '<w:sectPr><w:headerReference w:type="default" r:id="rIdH"/><w:footerReference w:type="default" r:id="rIdF"/>'
        '<w:pgSz w:w="11906" w:h="16838"/></w:sectPr>',
    ])
    files = {
        "[Content_Types].xml": HEAD + '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
            '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
            '<Default Extension="xml" ContentType="application/xml"/><Default Extension="jpeg" ContentType="image/jpeg"/>'
            '<Default Extension="png" ContentType="image/png"/>'
            '<Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>'
            '<Override PartName="/word/header1.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.header+xml"/>'
            '<Override PartName="/word/footer1.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.footer+xml"/>'
            '<Override PartName="/word/footnotes.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml"/>'
            '<Override PartName="/word/comments.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.comments+xml"/>'
            '<Override PartName="/docProps/core.xml" ContentType="application/vnd.openxmlformats-package.core-properties+xml"/>'
            '<Override PartName="/docProps/thumbnail.jpeg" ContentType="image/jpeg"/>'
            '</Types>',
        "_rels/.rels": HEAD + '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/>'
            '<Relationship Id="rId2" Type="http://schemas.openxmlformats.org/package/2006/relationships/metadata/core-properties" Target="docProps/core.xml"/>'
            '<Relationship Id="rId3" Type="http://schemas.openxmlformats.org/package/2006/relationships/metadata/thumbnail" Target="docProps/thumbnail.jpeg"/>'
            '</Relationships>',
        "word/_rels/document.xml.rels": HEAD + '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            '<Relationship Id="rIdH" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/header" Target="header1.xml"/>'
            '<Relationship Id="rIdF" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/footer" Target="footer1.xml"/>'
            '<Relationship Id="rIdN" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/footnotes" Target="footnotes.xml"/>'
            '<Relationship Id="rIdC" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/comments" Target="comments.xml"/>'
            '<Relationship Id="rIdI" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image" Target="media/image1.png"/>'
            '<Relationship Id="rIdMail" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink" Target="mailto:julien.moreau@exemple.fr" TargetMode="External"/>'
            '</Relationships>',
        "word/document.xml": HEAD + f"<w:document {NS}><w:body>{body}</w:body></w:document>",
        "word/header1.xml": HEAD + f"<w:hdr {NS}>" + para(run("Cabinet Dubreuil · Dossier "), run("Moreau", B)) + "</w:hdr>",
        "word/footer1.xml": HEAD + f"<w:ftr {NS}>" + para(run("Julien Moreau / Karim Hamidi, page "),
            '<w:fldSimple w:instr=" PAGE "><w:r><w:t>1</w:t></w:r></w:fldSimple>') + "</w:ftr>",
        "word/footnotes.xml": HEAD + f"<w:footnotes {NS}>"
            '<w:footnote w:type="separator" w:id="-1"><w:p><w:r><w:separator/></w:r></w:p></w:footnote>'
            '<w:footnote w:id="1">' + para(run("Voir le courrier de Sophie Lambert.")) + "</w:footnote></w:footnotes>",
        "word/comments.xml": HEAD + f"<w:comments {NS}>"
            '<w:comment w:id="0" w:author="Claire Dubreuil" w:initials="CD" w:date="2026-09-01T10:00:00Z">'
            + para(run("Vérifier avec Julien Moreau.")) + "</w:comment></w:comments>",
        "docProps/core.xml": HEAD + '<cp:coreProperties xmlns:cp="http://schemas.openxmlformats.org/package/2006/metadata/core-properties" '
            'xmlns:dc="http://purl.org/dc/elements/1.1/" xmlns:dcterms="http://purl.org/dc/terms/" '
            'xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance">'
            '<dc:title>Contrat Julien Moreau</dc:title><dc:creator>Claire Dubreuil</dc:creator>'
            '<cp:lastModifiedBy>Claire Dubreuil</cp:lastModifiedBy>'
            '<dcterms:created xsi:type="dcterms:W3CDTF">2026-09-01T10:00:00Z</dcterms:created></cp:coreProperties>',
        "docProps/thumbnail.jpeg": b"\xff\xd8\xff\xe0fake thumbnail\xff\xd9",
        "word/media/image1.png": b"\x89PNG\r\n\x1a\nfake image",
    }
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        for name, data in files.items():
            z.writestr(name, data)
    return out.getvalue()


def with_python_docx():
    """The same kind of content saved by python-docx (Word's own default template), or None
    when python-docx is not installed: checks the rewrite on a package Word itself shaped."""
    try:
        import docx
    except ImportError:
        return None
    d = docx.Document()
    d.add_heading("Dossier Julien Moreau", level=1)
    p = d.add_paragraph("Maître Claire Dubreuil pour ")
    p.add_run("Julien Mo").bold = True
    p.add_run("reau").italic = True
    p.add_run(", contre Karim Hamidi.")
    d.add_paragraph("Premier point : courrier de Karim Hamidi.", style="List Number")
    t = d.add_table(rows=2, cols=3)
    for c, v in zip(t.rows[0].cells, ("Nom", "Prénom", "Ville")):
        c.text = v
    for c, v in zip(t.rows[1].cells, ("MOREAU", "Julien", "Lyon")):
        c.text = v
    d.sections[0].header.paragraphs[0].text = "Dossier Moreau"
    d.sections[0].footer.paragraphs[0].text = "Julien Moreau, confidentiel"
    d.core_properties.author = "Claire Dubreuil"
    d.core_properties.title = "Contrat Julien Moreau"
    out = io.BytesIO()
    d.save(out)
    return out.getvalue()
