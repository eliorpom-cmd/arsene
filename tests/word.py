"""A Word document protected and restored in place, without the engine: a fixed table of
invented names stands for the detector, so this runs anywhere in a second (Linux included).
The same path with the real engine: tests/table.py. Usage: python tests/word.py"""
import io, pathlib, re, sys, zipfile
ROOT = pathlib.Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))
sys.path.insert(0, str(ROOT / "tests"))
import docxfile
from docxgen import contract, with_python_docx
from lxml import etree

PAIRS = {"julien.moreau@exemple.fr": "aubin.godet@exemple.fr", "Julien Moreau": "Aubin Godet",
         "Karim Hamidi": "Nabil Kaci", "Sophie Lambert": "Odile Perrot", "Claire Dubreuil": "Maud Fontaine",
         "14 rue des Lilas": "8 rue du Four", "14 août 1979": "3 août 1979", "14/08/1979": "03/08/1979",
         "Julien": "Aubin", "Moreau": "Godet", "MOREAU": "GODET", "Karim": "Nabil", "Hamidi": "Kaci",
         "HAMIDI": "KACI", "Dubreuil": "Fontaine"}
REAL = ["Moreau", "MOREAU", "moreau", "Hamidi", "HAMIDI", "Lambert", "Dubreuil", "Lilas", "14/08/1979"]
ok = total = 0


def check(name, cond):
    global ok, total
    total += 1; ok += bool(cond)
    print(("ok   " if cond else "FAIL ") + name)


def mapper(pairs):
    rx = re.compile("|".join(r"(?<!\w)" + re.escape(k) + r"(?!\w)" for k in sorted(pairs, key=len, reverse=True)))
    return lambda t: rx.sub(lambda m: pairs[m.group(0)], t)


FWD, BACK = mapper(PAIRS), mapper({v: k for k, v in PAIRS.items()})


def parts(data):
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        return {n: z.read(n) for n in z.namelist()}


def canon(xml):
    """The XML as Word reads it: attribute order, quotes and xml:space set by us aside."""
    root = etree.fromstring(xml)
    for el in root.iter():
        if isinstance(el.tag, str):
            el.attrib.pop("{http://www.w3.org/XML/1998/namespace}space", None)
    return etree.tostring(root, method="c14n")


def runs(data, part="word/document.xml"):
    """(formatting, text) of every run holding text, in order."""
    root = etree.fromstring(parts(data)[part])
    W = docxfile.W
    out = []
    for r in root.iter(f"{{{W}}}r"):
        t = "".join(x.text or "" for x in r.iter(f"{{{W}}}t"))
        if t:
            rpr = r.find(f"{{{W}}}rPr")
            fmt = "".join(sorted(etree.QName(c).localname for c in rpr)) if rpr is not None else ""
            out.append((fmt, t))
    return out


def leaks(data):
    found = set()
    for name, raw in parts(data).items():
        for v in REAL:
            if v.encode() in raw:
                found.add(f"{name}:{v}")
    return found


src = contract()
text = docxfile.text(src)
check("lecture : le nom coupé en deux runs se lit entier", "M. Julien Moreau, né le" in text)
check("lecture : tableau lu ligne par ligne, cellules séparées par des tabulations", "MOREAU\tJulien\t14/08/1979" in text.split("\n"))
check("lecture : saut de ligne = nouvelle ligne (une valeur ne le traverse pas)", "Adresse : 14 rue des Lilas" in text.split("\n"))
check("lecture : en-tête, pied, note, commentaire, zone de texte, suppression suivie",
      all(x in text for x in ("Dossier Moreau", "Karim Hamidi, page", "courrier de Sophie Lambert",
                              "Vérifier avec Julien", "Témoin : Sophie Lambert")) and "Julien Moreau" in text.split("\n"))
check("lecture : propriétés et lien mailto", "Contrat Julien Moreau" in text and "mailto:julien.moreau@exemple.fr" in text)

prot, report = docxfile.transform(src, FWD)
check("protection : aucun vrai nom dans aucune partie du fichier", not leaks(prot))
if leaks(prot):
    print("     ", sorted(leaks(prot)))
check("protection : le texte est celui que la fonction a rendu", docxfile.text(prot) == FWD(text))
r = runs(prot)
check("protection : nom coupé réparti sur les mêmes runs (gras « Go », italique « det »)",
      ("b", "M. Aubin Go") in r and ("i", "det") in r)
check("protection : mise en forme des autres runs intacte", ("b", "Kaci") in r and ("", " du 3 mars.") in r)
doc = etree.fromstring(parts(prot)["word/document.xml"])
W = docxfile.W
check("protection : titre, liste, tableau, tabulation, saut de ligne gardés",
      doc.find(f".//{{{W}}}pStyle").get(f"{{{W}}}val") == "Heading1"
      and len(doc.findall(f".//{{{W}}}numPr")) == 2 and len(doc.findall(f".//{{{W}}}tr")) == 3
      and len(doc.findall(f".//{{{W}}}tc")) == 9 and doc.find(f".//{{{W}}}br") is not None
      and len(doc.findall(f".//{{{W}}}r/{{{W}}}tab")) == 1)
check("protection : en-tête et pied réécrits, champ PAGE intact",
      runs(prot, "word/header1.xml")[-1] == ("b", "Godet")
      and b'w:instr=" PAGE "' in parts(prot)["word/footer1.xml"])
check("protection : auteur du commentaire et de la révision remplacés",
      b'w:author="Maud Fontaine"' in parts(prot)["word/comments.xml"]
      and b'w:author="Maud Fontaine"' in parts(prot)["word/document.xml"])
check("protection : miniature de la 1re page retirée (et ses renvois)",
      "docProps/thumbnail.jpeg" not in parts(prot) and b"thumbnail" not in parts(prot)["_rels/.rels"]
      and b"thumbnail" not in parts(prot)["[Content_Types].xml"])
check("protection : image signalée, laissée telle quelle",
      report["images"] == 1 and parts(prot)["word/media/image1.png"] == parts(src)["word/media/image1.png"]
      and docxfile.warnings(report))

# image metadata: EXIF/XMP/IPTC/comment of a JPEG and text chunks of a PNG go, pixels stay
import struct, zlib
def _chunk(kind, body):
    return struct.pack(">I", len(body)) + kind + body + struct.pack(">I", zlib.crc32(kind + body))
png = (b"\x89PNG\r\n\x1a\n" + _chunk(b"IHDR", struct.pack(">IIBBBBB", 1, 1, 8, 2, 0, 0, 0))
       + _chunk(b"tEXt", b"Author\x00Julien Moreau") + _chunk(b"IDAT", zlib.compress(b"\x00\xff\x00\x00"))
       + _chunk(b"IEND", b""))
seg = lambda m, body: b"\xff" + bytes([m]) + struct.pack(">H", len(body) + 2) + body
jpg = (b"\xff\xd8" + seg(0xE0, b"JFIF\x00\x01\x01") + seg(0xE1, b"Exif\x00\x00Julien Moreau GPS")
       + seg(0xFE, b"photo de Julien Moreau") + seg(0xDB, b"\x00" * 65) + b"\xff\xda\x00\x02PIXELS\xff\xd9")
p2, j2 = docxfile.strip_image_metadata(png), docxfile.strip_image_metadata(jpg)
check("images : métadonnées retirées (EXIF, commentaire, texte PNG), pixels identiques",
      b"Moreau" not in p2 and b"Moreau" not in j2 and b"IDAT" in p2 and p2.endswith(_chunk(b"IEND", b""))
      and j2.endswith(b"\xff\xda\x00\x02PIXELS\xff\xd9") and b"JFIF" in j2 and len(j2) == len(jpg) - 53)

back, _ = docxfile.transform(prot, BACK)
a, b = parts(src), parts(back)
same = all(canon(a[n]) == canon(b[n]) for n in b if n.endswith((".xml", ".rels")) and n not in ("_rels/.rels", "[Content_Types].xml"))
check("aller-retour : chaque partie XML identique à l'original", same)
check("aller-retour : images et autres fichiers identiques à l'octet",
      all(a[n] == b[n] for n in b if not n.endswith((".xml", ".rels"))))
check("aller-retour : nom coupé revenu sur ses deux runs d'origine", ("b", "M. Julien Mo") in runs(back) and ("i", "reau") in runs(back))
check("aller-retour : seule la miniature manque", set(a) - set(b) == {"docProps/thumbnail.jpeg"})

# a replacement longer than the name, over three runs, and back
three = contract().replace(b"<w:t>Hamidi</w:t>", b"<w:t>Ha</w:t></w:r><w:r><w:t>mi</w:t></w:r><w:r><w:rPr><w:i/></w:rPr><w:t>di</w:t>")
LONG = {**PAIRS, "Karim Hamidi": "Nabil Garnier-Benali", "Hamidi": "Garnier-Benali"}
long_fwd, long_back = mapper(LONG), mapper({v: k for k, v in LONG.items()})
p3, _ = docxfile.transform(three, long_fwd)
check("nom sur trois runs → faux nom plus long, texte exact", "Nabil Garnier-Benali du 3 mars" in docxfile.text(p3))
b3, _ = docxfile.transform(p3, long_back)
check("nom sur trois runs : aller-retour au texte exact", docxfile.units(b3) == docxfile.units(three))

# fail closed
try:
    docxfile.transform(src, lambda t: t.replace("Moreau", "Go\ndet"))
    check("échec fermé : une fonction qui casse les lignes est refusée", False)
except ValueError:
    check("échec fermé : une fonction qui casse les lignes est refusée", True)
try:
    docxfile.text(b"PK\x03\x04 not a zip")
    check("échec fermé : un faux .docx est refusé", False)
except ValueError:
    check("échec fermé : un faux .docx est refusé", True)

pd = with_python_docx()
if pd:
    import docx
    pp, _ = docxfile.transform(pd, FWD)
    check("python-docx : aucun vrai nom dans le fichier protégé", not leaks(pp))
    reopened = docx.Document(io.BytesIO(pp))
    check("python-docx : le fichier protégé se rouvre, mise en forme gardée",
          [(r.bold, r.italic, r.text) for r in reopened.paragraphs[1].runs][1:3] == [(True, None, "Aubin Go"), (None, True, "det")]
          and reopened.tables[0].cell(1, 0).text == "GODET" and reopened.core_properties.author == "Maud Fontaine")
    pb, _ = docxfile.transform(pp, BACK)
    check("python-docx : aller-retour identique",
          all(canon(parts(pd)[n]) == canon(parts(pb)[n]) for n in parts(pb)
              if n.endswith((".xml", ".rels")) and n not in ("_rels/.rels", "[Content_Types].xml")))
else:
    print("(python-docx absent : 3 vérifications sautées)")

print(f"{ok}/{total}")
sys.exit(ok != total)
